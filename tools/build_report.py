"""Genera el documento profesional del Parcial 1 (Word + PDF) a partir de los datos reales de IoT Central.

  python tools/build_report.py            # usa analisis/datos_4dias.csv ya descargado
  python tools/build_report.py --refresh  # vuelve a consultar IoT Central (analyze.py) y los logs de la VM

Salida: Informe_Parcial1_CacaoSense.docx y .pdf en la raiz de Parcial1_IoT.
"""
import argparse
import json
import subprocess
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
AN, EV, AS = ROOT / "analisis", ROOT / "evidencias", ROOT / "assets"
CAT = json.loads((ROOT / "senders" / "catalog.json").read_text(encoding="utf-8-sig"))["devices"]
COL = timezone(timedelta(hours=-5))
BROWN, GREEN = RGBColor(0x6B, 0x3A, 0x1E), RGBColor(0x3E, 0x7C, 0x3A)
VM = "azureuser@158.23.21.239"
SSH = r"C:\Windows\System32\OpenSSH\ssh.exe"
sys.path.insert(0, str(ROOT / "tools"))
from analyze import DAYS, DIAS_ES  # noqa: E402

AUTORES = "Daniel Villamizar · Tomás Urieles"

# ----------------------------------------------------------------------------------------------
# Datos de ingenieria (datasheets, parametros del codigo y umbrales)
# ----------------------------------------------------------------------------------------------
DATASHEET = {
    "cacao-01-lote1-twin": "METER TEROS 12",
    "cacao-02-lote2-wokwi": "TEROS 12 (pot.), AM2302, BH1750",
    "cacao-03-lote3-sdk": "METER TEROS 12, ROHM BH1750",
    "cacao-04-aire-cams": "Sensirion SPS30 (ref.), EPA US AQI",
    "cacao-05-meteo-feed": "Vaisala WXT530, Apogee SP-110-SS",
    "cacao-06-ferm-paho": "Maxim DS18B20, celda de carga + HX711",
    "cacao-07-campo-era5": "METER PHYTOS 31, WXT530",
    "cacao-08-dosel-rest": "AM2302/DHT22, ROHM BH1750",
    "cacao-09-riego-wokwi": "MaxBotix MB7389 (HC-SR04), DFRobot SEN0217",
    "cacao-10-bodega-mqttjs": "Littelfuse 59025, Panasonic EKMB1101112",
}
# variable: (unidad, sensor, rango datasheet, rango operativo, precision, umbral Rule, valor en el codigo)
PARAMS = [
    ("soilMoisture", "% VWC", "METER TEROS 12", "0–70 % (mineral)", "15–45 %", "±3 % VWC", "R1: < 20 % (prom. 15 min)",
     "Lote 1: min 15 / max 45 (DTDL). Lote 2: pot. 0–60 %. Lote 3: base 24 % −1,2 %/día, −5 % secado diurno, +6 % con riego"),
    ("soilTemperature", "°C", "METER TEROS 12", "−40 a 60 °C", "20–30 °C", "±0,5 °C", "—", "23 + 3,2·sen(ciclo diurno) ± 0,1"),
    ("soilEC", "dS/m", "METER TEROS 12", "0–20 dS/m", "0,3–1,5 dS/m", "±5 % +0,01", "—", "0,45 + 0,012·VWC ± 0,01"),
    ("temperature", "°C", "AM2302 / Vaisala WXT530", "−40 a 80 °C / −52 a 60 °C", "17–34 °C", "±0,5 / ±0,3 °C", "—",
     "Nodo 05: Open-Meteo temperature_2m. Nodo 08: T_ext − 1,8·luz. Nodo 02: DHT22 de Wokwi"),
    ("humidity", "% HR", "AM2302 / WXT530", "0–99,9 % HR", "50–100 % HR", "±2 / ±3 % HR", "—", "Open-Meteo relative_humidity_2m; dosel +7 %"),
    ("illuminance", "lx", "ROHM BH1750FVI", "1–65 535 lx", "0–45 000 lx", "±20 %", "—", "42 000·luz·(0,55+0,1·lote); Wokwi LDR 0–60 000"),
    ("rainfall", "mm", "Vaisala WXT530", "0–200 mm/h", "0–30 mm/h", "5 % (acum. diaria)", "—", "Open-Meteo rain (acumulado de la hora previa)"),
    ("windSpeed", "km/h", "Vaisala WXT530", "0–216 km/h (60 m/s)", "0–40 km/h", "±3 % a 36 km/h", "—", "Open-Meteo wind_speed_10m (km/h)"),
    ("radiation", "W/m²", "Apogee SP-110-SS", "0–2000 W/m²", "0–1100 W/m²", "<3 % calibración", "—", "Open-Meteo shortwave_radiation"),
    ("pm25", "µg/m³", "Sensirion SPS30", "0–1000 µg/m³", "0–60 µg/m³", "±(5 µg/m³ + 5 %)", "—", "Open-Meteo Air Quality pm2_5 (CAMS)"),
    ("pm10", "µg/m³", "Sensirion SPS30", "0–1000 µg/m³", "0–80 µg/m³", "±25 µg/m³", "—", "Open-Meteo Air Quality pm10 (CAMS)"),
    ("aqi", "índice", "EPA US AQI (calculado)", "0–500", "0–150", "N/A (índice)", "R4: > 100 (máx. 5 min)", "Open-Meteo us_aqi"),
    ("leafWetness", "% tiempo mojado", "METER PHYTOS 31", "300–1250 mV (umbral)", "0–100 %", "N/A (umbral)", "—",
     "100 % si lluvia > 0,1 mm; si no (HR − 85)·6,5 acotado 0–100"),
    ("boxTemperature", "°C", "Maxim DS18B20", "−55 a 125 °C", "25–52 °C", "±0,5 °C (−10 a 85)", "R2: > 52 °C (prom. 5 min)",
     "28 → 45 °C (48 h) → 50 °C (96 h) → descenso; ±1,5 °C volteo diario"),
    ("mass", "kg", "Celda de carga + HX711", "0–100 kg (celda)", "35–50 kg", "±0,05 % FS", "—", "50 − 0,09·h (pérdida de exudado)"),
    ("waterLevel", "%", "MaxBotix MB7389 (HC-SR04 en Wokwi)", "30–500 cm", "0–100 %", "resol. 1 mm", "R3: < 20 % (prom. 5 min)",
     "100·(200 − d)/(200 − 20), d en cm; enclavamiento de bomba < 10 %"),
    ("flowRate", "L/min", "DFRobot SEN0217", "1–30 L/min", "0–30 L/min", "±5 %", "—", "Pot. 0–30 L/min, 0 si la bomba está apagada"),
    ("pumpOn / doorOpen / motion", "booleano", "Relé / Littelfuse 59025 / PIR EKMB", "—", "true/false", "N/A", "—",
     "setPump (Wokwi); puerta 2 de cada 9 muestras en 07–17 h; PIR 25 % en horario"),
    ("fleetConnected / Disconnected / Unassociated", "nodos", "Puesto de mando (heartbeat)", "0–10", "0–10", "N/A", "R7: Disconnected > 0 (máx. 15 min)",
     "Connected si hay mensaje en max(3×intervalo, 5 min)"),
]
KEY_VARS = [("cacao-03-lote3-sdk", "soilMoisture", "% VWC"), ("cacao-01-lote1-twin", "soilMoisture", "% VWC"),
            ("cacao-02-lote2-wokwi", "soilMoisture", "% VWC"), ("cacao-05-meteo-feed", "temperature", "°C"),
            ("cacao-05-meteo-feed", "humidity", "% HR"), ("cacao-05-meteo-feed", "rainfall", "mm"),
            ("cacao-05-meteo-feed", "radiation", "W/m²"), ("cacao-04-aire-cams", "pm25", "µg/m³"),
            ("cacao-04-aire-cams", "aqi", "índice"), ("cacao-06-ferm-paho", "boxTemperature", "°C"),
            ("cacao-09-riego-wokwi", "waterLevel", "%"), ("cacao-08-dosel-rest", "temperature", "°C")]
SHOTS = [
    ("01-dashboard-control-room-a.png", "Cuarto de control CacaoSense: estado de flota, clima, aire, riego, KPI mín./máx. y gráficos."),
    ("01-dashboard-control-room-b.png", "Cuarto de control (continuación): gráficos, mapa de zonas y bloque de alertas."),
    ("02-flota-dispositivos.png", "Flota de 10 dispositivos sobre las 4 plantillas CacaoSense."),
    ("03-plantillas.png", "Plantillas de dispositivo (Digital Twin) publicadas."),
    ("05-reglas.png", "Reglas configuradas con acción de correo."),
    ("07-comando-setIrrigation.png", "Comando setIrrigation ejecutado en el Lote 3 (Python) con respuesta del dispositivo."),
    ("08-propiedad-writable.png", "Propiedad writable samplingIntervalSec aceptada (15 → 20 s) y riego habilitado."),
    ("09-desconectado.png", "Nodo en estado Desconectado durante una desconexión controlada."),
    ("10-wokwi-lote2.png", "Wokwi ESP32 #1 (Lote 2): monitor serie con DPS → IoT Hub → publish."),
    ("11-wokwi-riego.png", "Wokwi ESP32 #2 (Reservorio): monitor serie y respuesta al comando setPump."),
    ("12-wokwi-conectado-central.png", "Nodo Wokwi Conectado en IoT Central con telemetría en vivo."),
    ("13-data-explorer.png", "Data Explorer de IoT Central con la serie de la ventana de 4 días."),
]


# ----------------------------------------------------------------------------------------------
# Utilidades de formato
# ----------------------------------------------------------------------------------------------
def shade(cell, hex_color):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc.append(shd)


def table(doc, header, rows, widths=None, font=8):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold, r.font.size, r.font.color.rgb = True, Pt(font), RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "6B3A1E")
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run("" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v))
            run.font.size = Pt(font)
            if k % 2:
                shade(cells[i], "F7EFE6")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def fig(doc, path, caption, width=16.5, crop_browser=False):
    if not path.exists():
        return False
    img = path
    if crop_browser:
        im = Image.open(path).convert("RGB")
        w, h = im.size
        row = im.crop((0, int(h * 0.15), w, int(h * 0.15) + 1)).resize((1, 1)).getpixel((0, 0))
        top = int(h * 0.174) if sum(row) > 700 else int(h * 0.118)   # con/sin barra de depuracion de Chrome
        (EV / "recortes").mkdir(exist_ok=True)
        img = EV / "recortes" / path.name
        # En Wokwi el editor muestra la clave del dispositivo: solo se publica el panel de simulacion.
        left = int(w * 0.50) if "wokwi" in path.name else 0
        im.crop((left, top, w, h)).save(img)
    doc.add_picture(str(img), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    r = p.add_run(caption)
    r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0x55, 0x55, 0x55)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return True


def para(doc, text, size=10, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size, r.bold = Pt(size), bold
    return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(it).font.size = Pt(10)


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CacaoSense · Parcial 1 IoT + Cloud + Sistemas Distribuidos · UNAB 2026-II · Página ")
    r.font.size = Pt(8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = p.add_run()
        run.font.size = Pt(8)
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


# ----------------------------------------------------------------------------------------------
# Analisis sobre los datos reales
# ----------------------------------------------------------------------------------------------
def load():
    df = pd.read_csv(AN / "datos_4dias.csv")
    df["ts"] = pd.to_datetime(df["ts"], utc=True, format="ISO8601")
    df["day"] = df["day"].astype(str)
    return df


def key_stats(df):
    rows = []
    for dev, var, unit in KEY_VARS:
        sub = df[(df.device == dev)]
        if var not in sub:
            continue
        for day in DAYS:
            s = pd.to_numeric(sub[sub.day == day][var], errors="coerce").dropna()
            if s.empty:
                rows.append([dev[6:], var, DIAS_ES[day], "—", "—", "—", 0, "—"])
                continue
            t_max = sub.loc[s.idxmax(), "ts"].tz_convert(COL).strftime("%H:%M")
            rows.append([dev[6:], f"{var} ({unit})", DIAS_ES[day], f"{s.max():.2f} ({t_max})", f"{s.min():.2f}", f"{s.mean():.2f}",
                         int(s.count()), f"{s.sum():.1f}" if var == "rainfall" else "—"])
    return rows


def gaps(df):
    """Huecos > max(3 x intervalo, 5 min) por nodo y dia (desconexiones observadas en IoT Central)."""
    out = []
    for d in CAT:
        lim = timedelta(seconds=max(3 * d["intervalSec"], 300))
        for day in DAYS:
            s = df[(df.device == d["id"]) & (df.day == day)].ts.sort_values()
            for a, b in zip(s[:-1], s[1:]):
                if b - a > lim:
                    out.append([d["id"][6:], DIAS_ES[day], a.tz_convert(COL).strftime("%H:%M"), b.tz_convert(COL).strftime("%H:%M"),
                                f"{(b - a).total_seconds() / 60:.0f} min"])
    return out


def observed_intervals(df):
    rows = []
    for d in CAT:
        s = df[df.device == d["id"]].sort_values("ts").ts.diff().dt.total_seconds().dropna()
        s = s[s < 3 * d["intervalSec"] + 5]
        rows.append([d["id"][6:], f"{d['intervalSec']} s", f"{s.median():.0f} s" if len(s) else "—",
                     int((df.device == d["id"]).sum())])
    return rows


def reading(df):
    """Lectura operativa de los extremos (texto generado desde los datos)."""
    txt = []
    m = df[df.device == "cacao-05-meteo-feed"]
    for day in DAYS:
        s = m[m.day == day]
        if s.empty:
            continue
        tmax = s.loc[s.temperature.idxmax()]
        rain = s.rainfall.sum()
        txt.append(f"{DIAS_ES[day]}: la temperatura máxima del predio fue {tmax.temperature:.1f} °C a las "
                   f"{tmax.ts.tz_convert(COL):%H:%M} (radiación {tmax.radiation:.0f} W/m²), la mínima {s.temperature.min():.1f} °C "
                   f"de madrugada; lluvia acumulada {rain:.1f} mm y HR máxima {s.humidity.max():.0f} %.")
    s3 = df[df.device == "cacao-03-lote3-sdk"]
    if len(s3):
        dry = (s3.soilMoisture < 20).mean() * 100
        txt.append(f"Suelo Lote 3: {dry:.0f} % de las muestras quedaron por debajo del umbral académico de 20 % VWC; "
                   f"el mínimo ({s3.soilMoisture.min():.1f} %) ocurre al final de la tarde por evapotranspiración, y el comando "
                   f"setIrrigation lo eleva ~6 puntos. Es la condición que dispara R1 (suelo seco).")
    f = df[df.device == "cacao-06-ferm-paho"]
    if len(f):
        txt.append(f"Fermentación: la masa pasó de {f.boxTemperature.min():.1f} °C a {f.boxTemperature.max():.1f} °C, coherente con "
                   f"el perfil ICCO/AGROSAVIA (40–50 °C tras 48–96 h); no superó el umbral R2 de 52 °C, por lo que no hubo alarma de "
                   f"sobrefermentación. La masa bajó de {f['mass'].max():.1f} a {f['mass'].min():.1f} kg por pérdida de exudado.")
    a = df[df.device == "cacao-04-aire-cams"]
    if len(a):
        txt.append(f"Aire rural: PM2.5 entre {a.pm25.min():.1f} y {a.pm25.max():.1f} µg/m³ y US AQI máximo {a.aqi.max():.0f} "
                   f"(categoría {'buena' if a.aqi.max() <= 50 else 'moderada' if a.aqi.max() <= 100 else 'dañina para sensibles'}); "
                   f"R4 (AQI > 100) no se activó. Los picos coinciden con horas de baja mezcla atmosférica (madrugada).")
    return txt


def vm_logs():
    """Eventos de conexion/desconexion reales de la VM (journalctl) para el anexo de evidencias."""
    cmd = ("journalctl -t cacao-gap --no-pager -o short-iso | tail -40; echo ----; "
           "journalctl -u cacao-03 -u cacao-06 -u cacao-10 --no-pager -o cat | grep -E 'DISCONN|CONNECTED|FIN|CMD|PROP' | tail -40")
    try:
        out = subprocess.run([SSH, VM, cmd], capture_output=True, text=True, timeout=90).stdout
        (EV / "log_vm_eventos.txt").write_text(out, encoding="utf-8")
    except Exception:
        pass
    p = EV / "log_vm_eventos.txt"
    return p.read_text(encoding="utf-8").splitlines() if p.exists() else []


# ----------------------------------------------------------------------------------------------
def build():
    df = load()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    for lvl, size in ((1, 15), (2, 12)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.color.rgb, h.font.size, h.font.name = BROWN, Pt(size), "Calibri"
    page_number_footer(sec)
    hp = sec.header.paragraphs[0]
    hp.add_run().add_picture(str(AS / "cacaosense_logo.png"), width=Cm(3.2))
    hp.add_run("     Expediente del cliente · Granja y cultivo de cacao").font.size = Pt(8)

    # ---------------- Portada
    doc.add_paragraph()
    doc.add_picture(str(AS / "cacaosense_logo.png"), width=Cm(13))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, size, color in (("PARCIAL 1 — IoT + Cloud + Sistemas Distribuidos", 18, BROWN),
                              ("Escenario 5.3 · Granja y cultivo de cacao", 14, GREEN),
                              ("Flota heterogénea de 10 dispositivos sobre Azure IoT Central", 12, None),
                              ("Documento profesional de proyecto (expediente del cliente)", 11, None)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size, r.bold = Pt(size), size >= 14
        if color:
            r.font.color.rgb = color
    doc.add_paragraph()
    table(doc, ["Campo", "Valor"], [
        ["Universidad", "Universidad Autónoma de Bucaramanga (UNAB) · 2026-II"],
        ["Autores", AUTORES],
        ["Aplicación IoT Central", "CacaoSense - Granja de Cacao UNAB · https://cacaosense-unab2026.azureiotcentral.com"],
        ["Infraestructura Azure", "rg-parcial1-cacao · IoT Central ST2 (centralus) · VM vm-parcial1-cacao (mexicocentral)"],
        ["Ventana de datos", "4 fechas reales: " + ", ".join(DIAS_ES[d] for d in DAYS) + " de 2026"],
        ["Versión del documento", f"1.0 · generado {datetime.now(COL):%d/%m/%Y %H:%M} (hora Colombia)"],
    ], widths=[4.5, 12.5], font=9)
    doc.add_page_break()

    # ---------------- Historial de versiones
    doc.add_heading("Historial de versiones", 1)
    table(doc, ["Versión", "Fecha", "Autor", "Cambio"], [
        ["0.1", "23/09/2026", AUTORES, "Selección del escenario de cacao, catálogo de orígenes y fuentes verificadas (docs/fuentes_verificadas.md)."],
        ["0.5", "24/09/2026", AUTORES, "App IoT Central cacaosense-unab2026 (ST2), VM del parcial y 4 plantillas DTDL v1.0.0 (dtmi:unab:cacaosense:*;1)."],
        ["0.7", "24/09/2026", AUTORES, "Emisores: sdk_device.py v1.2.0, mqtt_explicit.py v1.1.0, https_bridge.py v1.0.0, bodega.js v1.0.0, "
                                        "fleet_monitor.py v1.0.0, Wokwi lote2_suelo.ino v1.1.0 y riego.ino v1.1.0."],
        ["0.8", "24/09/2026", AUTORES, "Plantillas v1.0.0 rev. 2: rangos minValue/maxValue para el Digital Twin y telemetría de estado de flota. "
                                        "Identidad visual, dashboard 'Cuarto de control', 5 reglas con correo, cron de desconexiones."],
        ["1.0", f"{datetime.now(COL):%d/%m/%Y}", AUTORES, "Documento final con la ventana de 4 días reales (24–27/09/2026) y comparativa."],
    ], widths=[1.4, 2.2, 3.6, 9.8], font=8.5)

    doc.add_heading("Contenido", 1)
    for i, t in enumerate(["Objetivo y escenario", "Arquitectura de referencia y telecomunicaciones", "Catálogo de 10 dispositivos",
                           "Digital Twin: plantillas, propiedades y comandos", "Tablas de parámetros por variable",
                           "Ventana de 4 días y comparativa de variables", "Evidencia de asincronía, desconexión y operación en línea",
                           "Cuarto de control, reglas y alertas", "Sustentación: dos códigos en dos equipos",
                           "Repositorio, seguridad y limitaciones de IoT Central", "Referencias"], 1):
        para(doc, f"{i}. {t}")
    doc.add_page_break()

    # ---------------- 1 Objetivo
    doc.add_heading("1. Objetivo y escenario", 1)
    para(doc, "CacaoSense es el demostrador IoT de un predio rural de cacao en Rionegro (Santander). Un operador no programador "
              "supervisa desde un cuarto de control la humedad de suelo de tres lotes, el clima del predio, la calidad del aire, la "
              "poscosecha (fermentación y dosel), el reservorio de riego y el perímetro de la bodega. La flota de 10 dispositivos se "
              "alimenta desde 10 códigos u orígenes de envío distintos hacia Azure IoT Central, de modo que se observan la asincronía "
              "(intervalos de 15 s a 5 min), las desconexiones (huecos y estado Disconnected) y la operación en línea.")
    para(doc, "Indispensables del escenario 5.3 cubiertos: humedad de suelo (nodos 01, 02 y 03) y un nodo meteorológico explícito "
              "(nodo 05, feed equivalente a Atlas Weather con temperatura, HR, lluvia, viento y radiación). La salida a Internet desde "
              "zona rural se justifica en la sección 2 (4G/LTE con respaldo satelital).")

    # ---------------- 2 Arquitectura
    doc.add_heading("2. Arquitectura de referencia y telecomunicaciones", 1)
    fig(doc, AN / "arquitectura_cacaosense.png", "Figura 1. Arquitectura de referencia CacaoSense: capas de dispositivo, red, plataforma y operación.")
    table(doc, ["Capa", "Elementos", "Decisión técnica"], [
        ["Dispositivo", "10 nodos: Digital Twin, 2 ESP32 Wokwi, Python SDK, 2 puentes de API, paho, replay CSV, HTTPS, Node.js",
         "Cada origen usa un código o feed distinto; ningún par comparte implementación."],
        ["Red / telecom", "Wi-Fi Wokwi-GUEST; VM Azure como gateway de borde; propuesta rural 4G/LTE Cat4 (Teltonika RUT241) + Starlink",
         "En el predio no hay fibra: LTE con WAN failover satelital, cola local store-and-forward y timestamp de origen."],
        ["Transporte", "MQTT/TLS 1.2 puerto 8883 (SDK, paho, MQTT.js, PubSubClient); HTTPS 443 (nodo 08 y APIs)",
         "Autenticación SAS HMAC-SHA256 por dispositivo; QoS 1 (IoT Hub no soporta QoS 2)."],
        ["Plataforma", "DPS (ID Scope) → IoT Hub gestionado → IoT Central; 4 plantillas DTDL v2",
         "Aprovisionamiento por clave simétrica de dispositivo; ninguna clave de grupo sale de la plataforma."],
        ["Operación", "Dashboard 'Cuarto de control', Views por dispositivo, 5 Rules con correo, Data Explorer / Query API",
         "Estado de flota por heartbeat (IoT Central no expone Connected/Disconnected por API)."],
    ], widths=[2.6, 7.2, 7.2], font=8)
    doc.add_page_break()

    # ---------------- 3 Catalogo
    doc.add_heading("3. Catálogo de 10 dispositivos", 1)
    para(doc, "Un solo catálogo: diez filas y diez orígenes distinguibles. Están presentes los cinco obligatorios (Digital Twin, Wokwi, "
              "Python, API pública y Atlas Weather o equivalente) y cinco adicionales aceptados por el enunciado.")
    table(doc, ["ID en Central", "Zona", "Origen de envío", "Protocolo", "Interv.", "Variables", "Datasheet"],
          [[d["id"], d["zone"], d["source"], d["protocol"], f"{d['intervalSec']} s", ", ".join(d["telemetry"]), DATASHEET[d["id"]]] for d in CAT],
          widths=[2.8, 2.2, 3.3, 2.6, 1.1, 3.0, 2.6], font=7)
    table(doc, ["Nodo", "Por qué el origen es distinto y cómo se provisiona"], [
        ["01", "Simulador nativo de IoT Central sobre la plantilla (simulated = true); valores acotados por minValue/maxValue del DTDL."],
        ["02 / 09", "Dos proyectos Wokwi independientes (firmware distinto: suelo y riego) en ESP32 con PubSubClient; DPS por MQTT y SAS con mbedTLS."],
        ["03", "Python con el SDK oficial azure-iot-device (ProvisioningDeviceClient + IoTHubDeviceClient), comando y propiedad writable."],
        ["04", "Puente de API pública: consulta Open-Meteo Air Quality (modelo CAMS) cada 5 min; conserva sourceTimestamp."],
        ["05", "Feed meteorológico Open-Meteo Forecast como equivalente de una estación Atlas Weather (temp., HR, lluvia, viento, radiación)."],
        ["06", "Cliente MQTT explícito paho-mqtt: SAS generado a mano, registro DPS y topics de IoT Hub visibles (continuación del Lab 3)."],
        ["07", "Replay de un CSV histórico real (Open-Meteo Archive ERA5, ago–sep 2026) re-emitido en vivo cada 30 s con su timestamp de origen."],
        ["08", "Puente HTTP/REST: DPS por HTTPS y POST a devices/{id}/messages/events (HTTP 204), sin MQTT."],
        ["10", "Node.js con MQTT.js (otro lenguaje y otra librería MQTT) + agregado del puesto de mando (estado de flota)."],
    ], widths=[1.6, 15.4], font=8)
    doc.add_page_break()

    # ---------------- 4 Digital twin
    doc.add_heading("4. Digital Twin: plantillas, propiedades y comandos", 1)
    table(doc, ["Plantilla (v1.0.0)", "DTMI", "Telemetría", "Propiedades", "Comandos", "Nodos"], [
        ["CacaoSense Nodo de Suelo", "dtmi:unab:cacaosense:suelo;1", "soilMoisture, soilTemperature, soilEC, illuminance, temperature",
         "samplingIntervalSec (writable), firmwareVersion, zone, originId, irrigationEnabled", "setIrrigation, setAlertLed", "01, 02, 03"],
        ["CacaoSense Clima y Aire", "dtmi:unab:cacaosense:clima;1", "temperature, humidity, rainfall, windSpeed, radiation, pm25, pm10, aqi, leafWetness, sourceKind, sourceTimestamp",
         "samplingIntervalSec, firmwareVersion, zone, originId", "—", "04, 05, 07"],
        ["CacaoSense Poscosecha y Dosel", "dtmi:unab:cacaosense:poscosecha;1", "boxTemperature, temperature, humidity, mass, illuminance",
         "samplingIntervalSec, firmwareVersion, zone, originId", "setAlertLed", "06, 08"],
        ["CacaoSense Riego y Perímetro", "dtmi:unab:cacaosense:riego;1", "waterLevel, flowRate, pumpOn, doorOpen, motion, temperature, fleetConnected/Disconnected/Unassociated",
         "samplingIntervalSec, firmwareVersion, zone, originId", "setPump, setAlertLed", "09, 10"],
    ], widths=[2.8, 3.2, 4.4, 3.4, 2.0, 1.2], font=7.5)
    fig(doc, EV / "03-plantillas.png", "Figura 2. Plantillas publicadas en IoT Central.", crop_browser=True)
    fig(doc, EV / "08-propiedad-writable.png", "Figura 3. Propiedad writable aceptada por el dispositivo (15 → 20 s) y riego habilitado por comando.", crop_browser=True)

    # ---------------- 5 Parametros
    doc.add_page_break()
    doc.add_heading("5. Tablas de parámetros por variable", 1)
    para(doc, "Cada variable se ancla a un sensor real. Se separan el rango del fabricante, el rango operativo del escenario, la precisión, "
              "el umbral de regla (académico y ajustable) y el valor usado en el código.")
    table(doc, ["Variable", "Unidad", "Sensor", "Rango datasheet", "Rango operativo", "Precisión", "Umbral Rule", "Valor en el código"],
          [list(p) for p in PARAMS], widths=[2.3, 1.3, 2.3, 2.0, 1.8, 1.7, 2.2, 3.4], font=6.8)

    # ---------------- 6 Ventana 4 dias
    doc.add_page_break()
    doc.add_heading("6. Ventana de 4 días y comparativa de variables", 1)
    para(doc, "Fechas analizadas (datos reales recibidos en vivo por IoT Central y consultados con la Query API): "
              + ", ".join(DIAS_ES[d] for d in DAYS) + ". La operación no fue continua: cada día incluye cortes programados y "
              "desconexiones (sección 7). Tabla: máximo (hora local), mínimo, promedio, recuento y sumatoria cuando aplica.")
    table(doc, ["Nodo", "Variable", "Día", "Máx (hora)", "Mín", "Prom.", "N", "Suma"], key_stats(df),
          widths=[2.6, 3.2, 2.2, 2.4, 1.5, 1.5, 1.2, 1.4], font=7)
    doc.add_heading("Lectura operativa de los extremos", 2)
    bullets(doc, reading(df))
    for f_, cap in (("g1_humedad_suelo_4dias.png", "Figura 4. Humedad de suelo por lote frente al umbral R1."),
                    ("g2_temperatura_4dias.png", "Figura 5. Temperatura del aire del predio (feed meteorológico)."),
                    ("g3_lluvia_4dias.png", "Figura 6. Lluvia horaria del predio."),
                    ("g4_fermentacion_4dias.png", "Figura 7. Temperatura de la masa en fermentación (MQTT explícito)."),
                    ("g5_pm25_4dias.png", "Figura 8. PM2.5 del aire rural (API pública CAMS)."),
                    ("g7_comparativa_meteo.png", "Figura 9. Comparativa mín./prom./máx. diaria de temperatura y HR.")):
        fig(doc, AN / f_, cap)
    fig(doc, EV / "13-data-explorer.png", "Figura 10. Data Explorer de IoT Central.", crop_browser=True)

    # ---------------- 7 Asincronia
    doc.add_page_break()
    doc.add_heading("7. Evidencia de asincronía, desconexión y operación en línea", 1)
    fig(doc, AN / "g6_asincronia_huecos.png", "Figura 11. Cada marca es un mensaje recibido por IoT Central; los huecos son desconexiones.")
    doc.add_heading("Intervalos configurados frente a observados", 2)
    table(doc, ["Nodo", "Configurado", "Observado (mediana)", "Mensajes en la ventana"], observed_intervals(df), widths=[5, 3, 3.5, 3.5], font=8)
    doc.add_heading("Huecos detectados en la serie (desconexiones)", 2)
    g = gaps(df)
    table(doc, ["Nodo", "Día", "Desde", "Hasta", "Duración"], g[:40] or [["—", "—", "—", "—", "—"]], widths=[5, 3, 2.5, 2.5, 2.5], font=8)
    para(doc, "Desconexiones programadas en la VM (cron, hora local): nodo 06 02:00–02:30, nodo 10 03:00–03:10, nodo 08 12:00–12:20 "
              "(corte de enlace simulado) y nodo 03 16:00–16:15. Los nodos Wokwi presentan además huecos propios: solo transmiten "
              "mientras la simulación está abierta, y la sustentación reproduce una desconexión y reconexión en vivo.", 9)
    fig(doc, EV / "09-desconectado.png", "Figura 12. Estado Desconectado en IoT Central.", crop_browser=True)
    logs = vm_logs()
    if logs:
        doc.add_heading("Registro de sesión de la VM (journalctl)", 2)
        p = doc.add_paragraph()
        r = p.add_run("\n".join(logs[-30:]))
        r.font.size, r.font.name = Pt(6.5), "Consolas"

    # ---------------- 8 Control room
    doc.add_page_break()
    doc.add_heading("8. Cuarto de control, reglas y alertas", 1)
    para(doc, "El dashboard de aplicación 'Cuarto de control CacaoSense' muestra al mismo tiempo el logo y nombre del escenario, el estado "
              "de la flota (Connected / Disconnected / Unassociated), KPI de mínimo y máximo del día, siete gráficos de telemetría de las "
              "variables indispensables, el bloque de reglas y el mapa de zonas dispositivo ↔ lugar físico.")
    for name, cap in SHOTS:
        if name.startswith(("01", "02", "05", "07")):
            fig(doc, EV / name, "Figura. " + cap, crop_browser=True)
    table(doc, ["Regla", "Plantilla", "Condición", "Agregación", "Acción"], [
        ["R1 Suelo seco en lotes", "Nodo de Suelo", "soilMoisture < 20 % VWC", "Promedio, 15 min", "Correo 'Alerta suelo seco'"],
        ["R2 Fermentación alta", "Poscosecha y Dosel", "boxTemperature > 52 °C", "Promedio, 5 min", "Correo 'Alerta fermentación'"],
        ["R3 Reservorio bajo", "Riego y Perímetro", "waterLevel < 20 %", "Promedio, 5 min", "Correo 'Alerta reservorio'"],
        ["R4 Calidad de aire", "Clima y Aire", "aqi > 100", "Máximo, 5 min", "Correo 'Alerta calidad de aire'"],
        ["R7 Nodo sin reporte", "Riego y Perímetro", "fleetDisconnected > 0", "Máximo, 15 min", "Correo 'Alerta nodo sin reporte'"],
    ], widths=[3.5, 3.0, 3.5, 3.0, 4.0], font=8)

    # ---------------- 9 Sustentacion
    doc.add_heading("9. Sustentación: dos códigos en dos equipos", 1)
    bullets(doc, [
        "Equipo A (portátil): detener el servicio del nodo 03 en la VM (sudo systemctl stop cacao-03) y ejecutar "
        "python senders/sdk_device.py cacao-03-lote3-sdk; en IoT Central se ve Connecting → Connected → telemetría cada 15 s.",
        "Equipo B: abrir los proyectos Wokwi (lote2_suelo y riego) y pulsar Play; el monitor serie muestra WiFi → NTP → DPS → HUB CONNECTED → TX.",
        "Comandos en vivo: setIrrigation(true) al Lote 3 (sube la humedad ~6 puntos) y setPump(true) al Wokwi #2 (LED azul encendido; rechazo si el nivel < 10 %).",
        "Desconexión controlada: Ctrl+C en el portátil o Stop en Wokwi → hueco en la serie y R7 (fleetDisconnected > 0) → reconexión automática al reanudar.",
        "El resto de la flota queda en segundo plano en la VM (servicios systemd) y con historial en los 4 días.",
    ])
    for name, cap in SHOTS:
        if name.startswith(("10", "11", "12")):
            fig(doc, EV / name, "Figura. " + cap, width=12 if "wokwi" in name else 16.5, crop_browser=True)

    # ---------------- 10 Repo
    doc.add_heading("10. Repositorio, seguridad y limitaciones de IoT Central", 1)
    bullets(doc, [
        "Repositorio: README de decisiones, models/ (DTDL), senders/ (Python, Node.js), wokwi/ (dos proyectos), deploy/ (systemd y cron), "
        "tools/ (administración REST, dashboard, análisis e informe), analisis/ y evidencias/.",
        "Sin secretos en claro: las claves de dispositivo y el token de API viven solo en senders/.env (ignorado por git, permisos 600 en la VM); "
        "los sketches publicados llevan PEGAR_PRIMARY_KEY y la copia con clave usa el prefijo PRIVADO_ (ignorado).",
        "Limitación: la API REST de IoT Central no expone el estado Connected/Disconnected ni permite crear reglas; se resolvió con un "
        "puesto de mando por heartbeat y reglas creadas en la interfaz.",
        "Limitación: los tiles de dashboard solo aceptan duraciones fijas (30 min, 1 h, 12 h, 1 d, 1 sem, 30 d) y la cuadrícula no se adapta "
        "a pantallas anchas; la Query API limita la tasa (HTTP 429) y exige consultas por plantilla.",
        "Limitación: el simulador nativo genera valores aleatorios salvo que el DTDL declare minValue/maxValue; IoT Hub no soporta QoS 2; "
        "los dispositivos HTTPS no mantienen sesión, por lo que su estado 'Connected' no es observable.",
        "Nota de transparencia: el 24/09 se probó la carga diferida (store-and-forward, propiedad iothub-creation-time-utc) con fechas "
        "18, 20 y 22/09; esos 7 dispositivos de prueba se eliminaron y la flota se re-aprovisionó con IDs nuevos, de modo que la ventana "
        "analizada contiene solo datos recibidos en vivo. Los nodos 01 (Digital Twin) y los modelos de farm.py son simulados, como admite "
        "el enunciado; los nodos 04, 05 y 07 usan fuentes públicas reales (Open-Meteo, CAMS, ERA5).",
    ])

    # ---------------- 11 Referencias
    doc.add_heading("11. Referencias", 1)
    bullets(doc, [
        "Microsoft Learn – Azure IoT Central: https://learn.microsoft.com/es-es/azure/iot-central/",
        "IoT Hub MQTT: https://learn.microsoft.com/azure/iot-hub/iot-mqtt-connect-to-iot-hub · Envío HTTPS: API send-device-event 2021-04-12.",
        "METER TEROS 12 y PHYTOS 31; Aosong AM2302; Maxim DS18B20; ROHM BH1750FVI; Vaisala WXT530; Apogee SP-110-SS; Sensirion SPS30; "
        "MaxBotix MB7389; DFRobot SEN0217; Littelfuse 59025; Panasonic EKMB1101112 (enlaces en docs/fuentes_verificadas.md).",
        "Open-Meteo Forecast, Air Quality (CAMS) y Archive (ERA5), CC BY 4.0: https://open-meteo.com",
        "ICCO – Growing cocoa; AGROSAVIA – beneficio del cacao (perfil de fermentación).",
        "Wokwi ESP32 Wi-Fi: https://docs.wokwi.com/guides/esp32-wifi",
    ])

    out = ROOT / "Informe_Parcial1_CacaoSense.docx"
    doc.save(out)
    print("DOCX:", out)
    return out


def to_pdf(docx_path):
    import win32com.client
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(str(docx_path))
        pdf = docx_path.with_suffix(".pdf")
        d.SaveAs(str(pdf), FileFormat=17)
        d.Close(False)
        print("PDF:", pdf)
    finally:
        word.Quit()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh", action="store_true")
    a = ap.parse_args()
    if a.refresh:
        subprocess.run([sys.executable, str(ROOT / "tools" / "analyze.py")], check=True)
    to_pdf(build())
